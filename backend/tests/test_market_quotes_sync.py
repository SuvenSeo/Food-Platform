from datetime import datetime, timezone

import pytest
from sqlalchemy import delete, select

from app.db.migrate import run_upgrade
from app.db.session import SessionLocal
from app.models.tables import MarketQuoteRecord
from app.services import market_quotes
from app.scrapers.dcs import _parse_dcs_table_rows
from app.scrapers.doa import parse_doa_item_rows
from app.scrapers.fisheries import parse_fisheries_workbook
from app.scrapers.harti import _parse_harti_price_tables, _parse_price_range


def _reset_market_quotes() -> None:
    run_upgrade()
    with SessionLocal() as db:
        db.execute(delete(MarketQuoteRecord))
        db.commit()


def _seed_quote(item_name: str = "Tomato") -> None:
    with SessionLocal() as db:
        db.add(
            MarketQuoteRecord(
                district="Colombo",
                market_name="Pettah",
                item_name=item_name,
                category="vegetables",
                unit="kg",
                price_lkr=320.0,
                source="seed",
                quoted_at=datetime.now(timezone.utc),
                notes=None,
            )
        )
        db.commit()


class _FakeResponse:
    def __init__(self, payload: object) -> None:
        self._payload = payload

    def raise_for_status(self) -> None:
        return None

    def json(self) -> object:
        return self._payload


def test_ingest_market_quotes_from_url_replaces_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    _reset_market_quotes()
    _seed_quote("Old Tomato")

    payload = [
        {
            "district": "Kandy",
            "market_name": "Kandy Central",
            "item_name": "Carrot",
            "category": "vegetables",
            "unit": "kg",
            "price_lkr": 450,
            "source": "remote-feed",
            "quoted_at": "2026-05-12T03:00:00Z",
            "notes": "remote sample",
        }
    ]
    monkeypatch.setattr(market_quotes.httpx, "get", lambda *_args, **_kwargs: _FakeResponse(payload))

    result = market_quotes.ingest_market_quotes_from_url(
        "https://example.com/quotes.json",
        timeout_seconds=5,
        payload_format="json",
    )

    assert result["market_quotes_count"] == 1
    with SessionLocal() as db:
        rows = db.scalars(select(MarketQuoteRecord).order_by(MarketQuoteRecord.id.asc())).all()
    assert len(rows) == 1
    assert rows[0].item_name == "Carrot"
    assert rows[0].source == "remote-feed"


def test_ingest_market_quotes_from_url_keeps_existing_rows_on_parse_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    _reset_market_quotes()
    _seed_quote("Tomato")

    invalid_payload = [{"district": "Colombo"}]
    monkeypatch.setattr(market_quotes.httpx, "get", lambda *_args, **_kwargs: _FakeResponse(invalid_payload))

    with pytest.raises(ValueError, match="missing required fields"):
        market_quotes.ingest_market_quotes_from_url(
            "https://example.com/invalid.json",
            timeout_seconds=5,
            payload_format="json",
        )

    with SessionLocal() as db:
        rows = db.scalars(select(MarketQuoteRecord).order_by(MarketQuoteRecord.id.asc())).all()
    assert len(rows) == 1
    assert rows[0].item_name == "Tomato"


def test_parse_doa_item_rows_emits_wholesale_and_retail_quotes() -> None:
    rows = [
        {
            "Date": "2026-05-18",
            "Item": "Beans",
            "Unit": "kg",
            "Pettah_Today_Wholesale": "650.000",
            "Dambulla_Today_Wholesale": "500.000",
            "Pettah_Today_Retail": "700.000",
            "Dambulla_Today_Retail": "530.000",
        }
    ]

    quotes = parse_doa_item_rows(rows, item="Beans")

    assert len(quotes) == 4
    assert quotes[0] == {
        "district": "Colombo",
        "market_name": "Manning Market (DOA)",
        "item_name": "Beans",
        "category": "vegetables",
        "unit": "kg",
        "price_lkr": 650.0,
        "source": "doa",
        "quoted_at": "2026-05-18T00:00:00+00:00",
        "notes": "DOA SHEP wholesale vegetable price.",
    }
    assert {quote["market_name"] for quote in quotes} == {
        "Manning Market (DOA)",
        "Dambulla Economic Centre (DOA)",
    }
    assert {quote["price_lkr"] for quote in quotes} == {650.0, 500.0, 700.0, 530.0}


def test_parse_doa_item_rows_accepts_wrapped_data_payload() -> None:
    payload = {
        "data": [
            {
                "Date": "2026-05-18",
                "Item": "Carrot",
                "Unit": "kg",
                "Pettah_Today_Wholesale": "250.000",
            }
        ]
    }

    quotes = parse_doa_item_rows(payload, item="Carrot")

    assert len(quotes) == 1
    assert quotes[0]["item_name"] == "Carrot"
    assert quotes[0]["price_lkr"] == 250.0


def test_parse_dcs_table_rows_uses_current_week_average_price() -> None:
    quoted_at = datetime(2026, 5, 14, tzinfo=timezone.utc)
    rows = [
        ["Item", "Unit", "Avg.Price (Rs.cts.)", "% Change", "Price range"],
        ["Ash Plantain", "1 kg.", "271.97 274.81 281.43", "3.5% 2.4%", "240.00 320.00"],
        ["Gotukola", "Bunch", "60.78 57.50 55.56", "-8.6% -3.4%", "40.00 80.00"],
    ]

    quotes = _parse_dcs_table_rows(rows, quoted_at)

    assert quotes == [
        {
            "district": "Colombo",
            "market_name": "Colombo District (DCS)",
            "item_name": "Ash Plantain",
            "category": "vegetables",
            "unit": "kg",
            "price_lkr": 281.43,
            "source": "dcs",
            "quoted_at": "2026-05-14T00:00:00+00:00",
            "notes": "DCS weekly open market retail price survey, Colombo District.",
        },
        {
            "district": "Colombo",
            "market_name": "Colombo District (DCS)",
            "item_name": "Gotukola",
            "category": "vegetables",
            "unit": "bunch",
            "price_lkr": 55.56,
            "source": "dcs",
            "quoted_at": "2026-05-14T00:00:00+00:00",
            "notes": "DCS weekly open market retail price survey, Colombo District.",
        },
    ]


def test_parse_harti_price_range_averages_ranges_and_rejects_blanks() -> None:
    assert _parse_price_range("550 - 650") == 600.0
    assert _parse_price_range("350") == 350.0
    assert _parse_price_range("-") is None


def test_parse_harti_price_tables_emits_market_quotes() -> None:
    tables = [
        [
            ["", "2026.05.18", "2026.05.18"],
            ["Variety", "Peliyagoda Market", "Kandy Market"],
            ["Up Country Vegetable", "", ""],
            ["Beans", "550 - 650", "500 - 550"],
            ["Banana", "", ""],
            ["Anamalu (Rs/Fruits)", "20 - 29", "-"],
        ]
    ]

    quotes = _parse_harti_price_tables(
        tables,
        pdf_url="https://www.harti.gov.lk/assets/pdf/food_price/daily/eng/2026/May/Vegetable%20Pricenew%20ex1(2026.05.18).pdf",
    )

    assert quotes == [
        {
            "district": "Colombo",
            "market_name": "Peliyagoda Market (HARTI)",
            "item_name": "Beans",
            "category": "vegetables",
            "unit": "kg",
            "price_lkr": 600.0,
            "source": "harti",
            "quoted_at": "2026-05-18T00:00:00+00:00",
            "notes": "HARTI daily food commodities bulletin range: 550 - 650.",
        },
        {
            "district": "Kandy",
            "market_name": "Kandy Market (HARTI)",
            "item_name": "Beans",
            "category": "vegetables",
            "unit": "kg",
            "price_lkr": 525.0,
            "source": "harti",
            "quoted_at": "2026-05-18T00:00:00+00:00",
            "notes": "HARTI daily food commodities bulletin range: 500 - 550.",
        },
        {
            "district": "Colombo",
            "market_name": "Peliyagoda Market (HARTI)",
            "item_name": "Anamalu",
            "category": "fruits",
            "unit": "piece",
            "price_lkr": 24.5,
            "source": "harti",
            "quoted_at": "2026-05-18T00:00:00+00:00",
            "notes": "HARTI daily food commodities bulletin range: 20 - 29.",
        },
    ]


def test_parse_fisheries_workbook_emits_wholesale_and_retail_quotes() -> None:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    wholesale = workbook.active
    wholesale.title = "Wholesale"
    wholesale.append(["Table 1 : Change in Wholesale Prices at Peliyagoda Fish Market (Rs/Kg)"])
    wholesale.append(["Variety", None, None, 2025, 2026, None])
    wholesale.append(["Sinhala Name", None, "Common Name", "2nd week of Apr.", "1st week of Apr.", "2nd week of Apr."])
    wholesale.append([1, "seer-si", "Seer (Ni-L)", 1660, 2066.67, 2040])

    retail = workbook.create_sheet("Retail")
    retail.append(["Table 2: Change in Consumer Prices at Selected Markets - (Rs/Kg)"])
    retail.append(["Variety", None, None, 2025, 2026, None])
    retail.append(["Sinhala Name", None, "Common Name", "2nd week of Apr.", "1st week of Apr.", "2nd week of Apr."])
    retail.append([1, "skipjack-si", "Skipjack tuna", 1510, 1890, 1843.3])

    stream = BytesIO()
    workbook.save(stream)

    quotes = parse_fisheries_workbook(stream.getvalue(), label="2nd Week April 2026")

    assert len(quotes) == 2
    assert quotes[0]["source"] == "fisheries"
    assert quotes[0]["market_name"] == "Peliyagoda Fish Market (Fisheries)"
    assert quotes[0]["item_name"] == "Seer (Ni-L)"
    assert quotes[0]["category"] == "fish"
    assert quotes[0]["price_lkr"] == 2040.0
    assert quotes[0]["quoted_at"].startswith("2026-04-08")
    assert quotes[1]["market_name"] == "Selected Retail Markets (Fisheries)"
